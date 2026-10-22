import unittest
from decimal import Decimal
from unittest.mock import patch
from uuid import uuid4
from sqlalchemy import select, func
import test_purchase_requests as fixtures
from services import purchase_requests as pr
from services.erp_service import cancel_purchase_order, sync_legacy_purchase_order
from models import BomComponent, SKU, Movement, StockBalance
from purchase_request_export import build_workbook, COLUMNS
from openpyxl import load_workbook


class MultiplePurchaseOrderTests(unittest.TestCase):
    setUpBase = fixtures.PurchaseRequestTests.setUpBase
    setUp = fixtures.PurchaseRequestTests.setUp
    tearDown = fixtures.PurchaseRequestTests.tearDown
    request = fixtures.PurchaseRequestTests.request
    order = fixtures.PurchaseRequestTests.order

    def allocate(self, row, pairs, user=None):
        result = pr.transition(self.db, row['id'], dict(action='ALOCAR_PEDIDO',
            version=row['version'], reason='Dividir atendimento entre fornecedores',
            allocations=[dict(purchase_order_id=order['id'], quantity=qty) for order, qty in pairs]),
            user or self.buyer)['request']
        self.db.commit()
        return result

    def test_two_orders_cover_one_request_without_duplicate_rows_or_stock(self):
        a, b = self.order([], qty='4'), self.order([], qty='6')
        row = self.request(quantity='10')
        result = self.allocate(row, [(a, '4'), (b, '6')])
        self.assertEqual('CONCLUIDA', result['status'])
        self.assertEqual(Decimal(10), Decimal(result['allocated_quantity']))
        self.assertEqual(Decimal(0), Decimal(result['remaining_quantity']))
        self.assertEqual({a['id'], b['id']}, {o['purchase_order_id'] for o in result['purchase_orders']})
        self.assertEqual(1, pr.listing(self.db, {})['total'])
        self.assertEqual(2, len(pr.history(self.db, row['id'])['events'][-1]['after_data']['purchase_orders']))
        self.assertEqual(0, self.db.query(Movement).count())
        self.assertEqual(0, self.db.query(StockBalance).count())

    def test_partial_persists_and_second_order_completes(self):
        a, b = self.order([], qty='4'), self.order([], qty='6')
        row = self.request(quantity='10')
        partial = self.allocate(row, [(a, '4')])
        self.assertEqual('EM_COMPRAS', partial['status'])
        self.assertIsNone(partial['completed_at'])
        self.assertIsNone(partial['purchase_order_id'])
        self.db.expire_all()
        listed = pr.listing(self.db, {}, self.buyer)['items'][0]
        self.assertEqual('6.000', listed['remaining_quantity'])
        result = self.allocate(listed, [(b, '6')])
        self.assertEqual('CONCLUIDA', result['status'])
        self.assertEqual(2, len(result['purchase_orders']))

    def test_partial_reservation_is_not_reused_by_other_request(self):
        a = self.order([], qty='5')
        row = self.request(quantity='10')
        self.allocate(row, [(a, '4')])
        second = self.request(quantity='3', sector='ADMINISTRATIVO')
        options = pr.existing_order_options(self.db, second['id'], self.buyer)['items']
        self.assertEqual('1.000', options[0]['available_quantity'])
        with self.assertRaisesRegex(ValueError, 'saldo pendente'):
            self.allocate(second, [(a, '2')])
        self.db.rollback()

    def test_sum_over_request_duplicate_order_and_invalid_quantities_fail(self):
        a, b = self.order([], qty='5'), self.order([], qty='5')
        row = self.request(quantity='4')
        for pairs in ([(a, '3'), (b, '2')], [(a, '1'), (a, '1')],
                      [(a, '0')], [(a, '-1')], [(a, 'NaN')], [(a, '1.0001')]):
            with self.subTest(pairs=pairs), self.assertRaises(ValueError):
                self.allocate(row, pairs)
            self.db.rollback()
        self.assertEqual([], pr.active_allocations(self.db, row['id']))

    def test_wrong_sku_or_closed_order_rolls_back_whole_batch(self):
        a, b = self.order([], qty='2'), self.order([], qty='3', sku='OTHER')
        row = self.request(quantity='5')
        with self.assertRaisesRegex(ValueError, 'saldo pendente'):
            self.allocate(row, [(a, '2'), (b, '3')])
        self.db.rollback()
        self.assertEqual([], pr.active_allocations(self.db, row['id']))
        self.assertEqual(1, len(pr.history(self.db, row['id'])['events']))

    def test_only_buyer_can_allocate_and_stale_version_is_rejected(self):
        a = self.order([], qty='5')
        row = self.request(quantity='5')
        with self.assertRaises(PermissionError):
            self.allocate(row, [(a, '2')], self.pcp)
        self.db.rollback()
        self.allocate(row, [(a, '2')])
        with self.assertRaisesRegex(ValueError, 'outro usuário'):
            self.allocate(row, [(a, '2')])
        self.db.rollback()

    def test_cancelling_one_order_keeps_other_allocation_and_reopens_only_missing(self):
        a, b = self.order([], qty='4'), self.order([], qty='6')
        row = self.request(quantity='10')
        self.allocate(row, [(a, '4'), (b, '6')])
        cancel_purchase_order(self.db, b['id'], self.buyer.username, 'Fornecedor cancelou')
        result = pr.history(self.db, row['id'])['request']
        self.assertEqual('EM_COMPRAS', result['status'])
        self.assertEqual('6.000', result['remaining_quantity'])
        self.assertEqual([a['id']], [o['purchase_order_id'] for o in result['purchase_orders']])
        self.assertEqual(2, len(pr.history(self.db, row['id'])['events'][-1]['before_data']['purchase_orders']))

    def test_any_linked_order_edit_cannot_reduce_allocated_quantity(self):
        akey, bkey = str(uuid4()), str(uuid4())
        a, b = self.order([], qty='4', key=akey), self.order([], qty='6', key=bkey)
        row = self.request(quantity='10')
        self.allocate(row, [(a, '4'), (b, '6')])
        data = fixtures.legacy.ErpSkuResolutionTest._payload(bkey, 'MAT-001', '5')
        with self.assertRaisesRegex(ValueError, 'quantidade'):
            sync_legacy_purchase_order(self.db, data, self.buyer.username)
        self.db.rollback()
        self.assertEqual('CONCLUIDA', pr.get(self.db, row['id'])['status'])

    def test_component_coverage_from_kit_is_proportional_in_multiple_orders(self):
        kit = SKU(sku='CJ-TESTE', descricao='Conjunto', unidade='CJ', grupo='30 - CONJUNTO', active=True)
        self.db.add(kit); self.db.flush()
        self.db.add(BomComponent(item_sku_id=kit.id, component_sku_id=self.active_sku.id, quantidade=Decimal('0.8')))
        self.db.commit()
        a, b = self.order([], qty='2', sku='CJ-TESTE'), self.order([], qty='1.4')
        row = self.request(quantity='3')
        result = self.allocate(row, [(a, '1.6'), (b, '1.4')])
        self.assertEqual('CONCLUIDA', result['status'])

    def test_search_secondary_order_and_excel_keep_all_allocations(self):
        a, b = self.order([], qty='4'), self.order([], qty='6')
        row = self.request(quantity='10')
        self.allocate(row, [(a, '4')])
        order_number = self.db.execute(select(pr.orders.c.numero_oc).where(
            pr.order_condition(self.db, pr.orders.c.id, a['id']))).scalar_one()
        self.assertEqual(1, pr.listing(self.db, {'q': order_number})['total'])
        data = pr.export_data(self.db, {}, self.buyer)
        wb = load_workbook(build_workbook(data, {}, self.buyer.username))
        headers = [c.value for c in wb['Solicitações'][4]]
        values = dict(zip(headers, next(wb['Solicitações'].iter_rows(min_row=5, values_only=True))))
        self.assertIn(order_number, values['Pedidos alocados / quantidades'])
        self.assertEqual(4, values['Quantidade alocada na solicitação'])
        self.assertEqual(6, values['Quantidade restante na solicitação'])

    def test_cancelling_request_releases_all_allocations_and_preserves_audit(self):
        a = self.order([], qty='4')
        row = self.allocate(self.request(quantity='10'), [(a, '4')])
        pr.transition(self.db, row['id'], dict(action='CANCELAR', version=row['version'], reason='Demanda cancelada'), self.buyer)
        self.db.commit()
        self.assertEqual([], pr.active_allocations(self.db, row['id']))
        self.assertEqual('4.000', pr.encode(pr._order_component_unallocated(self.db, a['id'], 'MAT-001')))

    def test_completed_request_can_be_redistributed_without_changing_original_completion(self):
        a, b = self.order([], qty='10'), self.order([], qty='6')
        row = self.allocate(self.request(quantity='10'), [(a, '10')])
        completed_at = row['completed_at']
        result = pr.transition(self.db, row['id'], dict(action='ALOCAR_PEDIDO', version=row['version'],
            replace_allocations=True, reason='Dividir fornecimento', allocations=[
                dict(purchase_order_id=a['id'], quantity='4'), dict(purchase_order_id=b['id'], quantity='6')]), self.buyer)['request']
        self.db.commit()
        self.assertEqual('CONCLUIDA', result['status'])
        self.assertEqual(completed_at, result['completed_at'])
        self.assertEqual(2, len(result['purchase_orders']))
        self.assertEqual(Decimal(10), sum(Decimal(o['quantity']) for o in result['purchase_orders']))

    def test_replacement_releases_removed_order_and_partial_reopens(self):
        a, b = self.order([], qty='4'), self.order([], qty='6')
        row = self.allocate(self.request(quantity='10'), [(a, '4'), (b, '6')])
        result = pr.transition(self.db, row['id'], dict(action='ALOCAR_PEDIDO', version=row['version'],
            replace_allocations=True, reason='Manter apenas pedido A', allocations=[
                dict(purchase_order_id=a['id'], quantity='4')]), self.buyer)['request']
        self.db.commit()
        self.assertEqual('EM_COMPRAS', result['status'])
        self.assertIsNone(result['completed_at'])
        self.assertEqual('6.000', result['remaining_quantity'])
        self.assertEqual(Decimal(6), pr._order_component_unallocated(self.db, b['id'], 'MAT-001'))

    def test_old_primary_link_fallback_is_visible_without_normalized_rows(self):
        a = self.order([], qty='10')
        row = self.allocate(self.request(quantity='10'), [(a, '10')])
        self.db.execute(pr.request_orders.delete().where(pr.request_orders.c.request_id == row['id']))
        self.db.commit()
        snapshot = pr.history(self.db, row['id'])['request']
        self.assertEqual(a['id'], snapshot['purchase_orders'][0]['purchase_order_id'])
        self.assertEqual(Decimal(0), pr._order_component_unallocated(self.db, a['id'], 'MAT-001'))

    def test_conversion_still_accepts_direct_request_for_kit_with_bom(self):
        kit = SKU(sku='CJ-DIRETO', descricao='Conjunto', unidade='CJ', grupo='30 - CONJUNTO', active=True)
        self.db.add(kit); self.db.flush()
        self.db.add(BomComponent(item_sku_id=kit.id, component_sku_id=self.active_sku.id, quantidade=2))
        self.db.commit()
        row = self.request(quantity='2', sku_codigo=kit.sku)
        order = self.order([row['id']], qty='2', sku=kit.sku)
        self.assertEqual('CONCLUIDA', pr.get(self.db, row['id'])['status'])
        self.assertEqual(order['id'], pr.active_allocations(self.db, row['id'])[0]['purchase_order_id'])

    def test_new_order_uses_only_remaining_quantity_after_partial_allocation(self):
        a = self.order([], qty='4')
        row = self.allocate(self.request(quantity='10'), [(a, '4')])
        row = pr.transition(self.db, row['id'], dict(action='CONVERTER_NOVA_COMPRA', version=row['version'], reason='Saldo em nova compra'), self.buyer)['request']
        self.db.commit()
        self.assertEqual('6.000', pr.prepare(self.db, [row['id']], self.buyer)['items'][0]['quantity'])
        self.order([row['id']], qty='6')
        current = pr.history(self.db, row['id'])['request']
        self.assertEqual('CONCLUIDA', current['status'])
        self.assertEqual(2, len(current['purchase_orders']))


if __name__ == '__main__':
    unittest.main()

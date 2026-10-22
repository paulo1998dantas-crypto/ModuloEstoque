import unittest
from decimal import Decimal
from unittest.mock import patch
import test_purchase_requests as fixtures
from models import SKU, BomComponent
from services import purchase_requests as pr
from services.purchase_request_components import explode_requests, load_graph
from purchase_request_export import build_workbook
from openpyxl import load_workbook


class ComponentProjectionTests(unittest.TestCase):
    def setUp(self):
        self.catalog = {code:dict(sku=code,descricao=code,unidade=unit,grupo=group,active=True)
            for code,group,unit in [("CJ","30","CJ"),("PP","20","PC"),
                                    ("MAT","10","M"),("KIT","30","CJ"),("OTHER","10","PC")]}
        self.row = dict(id="id-1",sku_codigo="CJ",descricao="Original",unidade="CJ",
            quantity="3",status="SOLICITADA",needed_at="2026-10-29",work_order_ids=["os-1","os-2"],
            anticipation_pending_quantity="5")
        self.bom = [dict(parent_sku="CJ",component_sku="PP",quantidade="2"),
                    dict(parent_sku="PP",component_sku="MAT",quantidade="0.8"),
                    dict(parent_sku="CJ",component_sku="MAT",quantidade="0.1"),
                    dict(parent_sku="CJ",component_sku="KIT",quantidade="2")]

    def test_nested_fractional_consumption_aggregates_same_leaf_only_within_request(self):
        original = dict(self.row)
        output = explode_requests([self.row],self.catalog,self.bom)
        self.assertEqual(["MAT","KIT"],[r["sku_codigo"] for r in output])
        self.assertEqual(Decimal("1.7"),output[0]["fator_bom"])
        self.assertEqual(Decimal("5.1"),output[0]["quantity"])
        self.assertEqual(Decimal("8.5"),output[0]["anticipation_pending_quantity"])
        self.assertEqual("M",output[0]["unidade"])
        self.assertEqual(Decimal("6"),output[1]["quantity"])
        self.assertEqual("CJ",output[0]["sku_solicitado"])
        for item in output:
            self.assertEqual(original["needed_at"],item["needed_at"])
            self.assertEqual(original["work_order_ids"],item["work_order_ids"])
        self.assertEqual(original,self.row)

    def test_no_bom_keeps_own_code_and_insumo_with_bom_does_not_explode(self):
        for code in ("KIT","PP","MAT"):
            row=dict(self.row,sku_codigo=code)
            output=explode_requests([row],self.catalog,[dict(parent_sku="MAT",component_sku="OTHER",quantidade=2)])
            self.assertEqual(code,output[0]["sku_codigo"])
            self.assertEqual("3",output[0]["quantity"])

    def test_different_requests_are_not_merged_even_for_same_sku(self):
        second=dict(self.row,id="id-2",needed_at="2026-11-15",quantity="2")
        output=explode_requests([self.row,second],self.catalog,self.bom)
        self.assertEqual(4,len(output))
        self.assertEqual({"id-1","id-2"},{r["id"] for r in output})

    def test_cycle_missing_child_and_invalid_consumption_are_errors(self):
        for bom in ([dict(parent_sku="CJ",component_sku="CJ",quantidade=1)],
                    [dict(parent_sku="CJ",component_sku="MISSING",quantidade=1)],
                    [dict(parent_sku="CJ",component_sku="MAT",quantidade=0)],
                    [dict(parent_sku="CJ",component_sku="MAT",quantidade="NaN")]):
            with self.subTest(bom=bom),self.assertRaises(ValueError):
                explode_requests([self.row],self.catalog,bom)

    def test_inactive_component_is_visible_and_flagged_not_silently_dropped(self):
        self.catalog["MAT"]["active"]=False
        output=explode_requests([self.row],self.catalog,self.bom)
        self.assertEqual("INATIVO",output[0]["cadastro_componente"])

    def test_online_load_reads_only_reachable_without_writes(self):
        from services import cadastro_supabase_service as source
        calls=[]
        def read(table, fields, order, filters=None):
            calls.append((table,filters))
            if table==source.REGISTRATIONS_TABLE:
                codes=[code for code in ("30000001","20000001","10000001") if '"'+code+'"' in str(filters)]
                return [dict(sku=code,descricao_primaria=code,unidade="PC",ativo=True) for code in codes]
            return [edge for edge in [
                dict(parent_sku="30000001",component_sku="20000001",quantidade="2"),
                dict(parent_sku="20000001",component_sku="10000001",quantidade="0.8")]
                if '"'+edge["parent_sku"]+'"' in str(filters)]
        with patch.object(source,"enabled",return_value=True),patch.object(source,"_all_rows",side_effect=read):
            catalog,bom=load_graph(None,{"30000001"})
        self.assertEqual(3,len(catalog));self.assertEqual(2,len(bom))
        self.assertTrue(all(filters for _,filters in calls))

    def test_cadastro_failure_is_not_disguised_as_missing_bom(self):
        from services import cadastro_supabase_service as source
        with patch.object(source,"enabled",return_value=True),patch.object(source,"_all_rows",
                side_effect=source.CadastroSupabaseError("Indisponível")):
            with self.assertRaisesRegex(ValueError,"exportação interrompida"):
                load_graph(None,{"30000001"})

    def test_filtered_pagination_keeps_filters_on_all_pages(self):
        from services import cadastro_supabase_service as source
        filters=[("parent_sku",'in.("30000001")')]
        with patch.object(source,"PAGE_SIZE",1),patch.object(source,"_request_rows",
                side_effect=[[{"id":1}],[{"id":2}],[]]) as request:
            rows=source._all_rows("cadastro_bom_componentes","id","id.asc",filters=filters)
        self.assertEqual([{"id":1},{"id":2}],rows)
        self.assertEqual([0,1,2],[call.args[3] for call in request.call_args_list])
        self.assertTrue(all(call.kwargs["filters"]==filters for call in request.call_args_list))


class ComponentExportIntegrationTests(unittest.TestCase):
    setUpBase=fixtures.PurchaseRequestTests.setUpBase
    setUp=fixtures.PurchaseRequestTests.setUp
    tearDown=fixtures.PurchaseRequestTests.tearDown
    request=fixtures.PurchaseRequestTests.request
    order=fixtures.PurchaseRequestTests.order

    def test_export_projects_bom_without_changing_requests_or_audit(self):
        kit=SKU(sku="30000001",descricao="Kit",unidade="CJ",grupo="30 - CONJUNTO",active=True)
        pp=SKU(sku="20000001",descricao="PP",unidade="PC",grupo="20 - PRODUTO PROCESSO",active=True)
        self.db.add_all([kit,pp]);self.db.flush()
        self.db.add_all([BomComponent(item_sku_id=kit.id,component_sku_id=pp.id,quantidade=2),
                        BomComponent(item_sku_id=pp.id,component_sku_id=self.active_sku.id,quantidade=Decimal("0.8"))])
        self.db.commit()
        row=self.request(sku_codigo=kit.sku,quantity="3")
        before=pr.history(self.db,row["id"])
        with patch("services.cadastro_supabase_service.enabled",return_value=False):
            data=pr.export_data(self.db,{},self.user)
        self.assertEqual(before,pr.history(self.db,row["id"]))
        self.assertEqual("30000001",data["history_items"][0]["sku_codigo"])
        self.assertEqual("MAT-001",data["component_items"][0]["sku_codigo"])
        self.assertEqual(Decimal("4.8"),data["component_items"][0]["quantity"])
        book=load_workbook(build_workbook(data))
        headers=[c.value for c in book["Solicitações"][4]]
        values=list(book["Solicitações"].iter_rows(min_row=5,values_only=True))[0]
        exported=dict(zip(headers,values))
        self.assertEqual("MAT-001",exported["SKU"])
        self.assertEqual("30000001",exported["SKU solicitado"])
        self.assertEqual(4.8,exported["Quantidade"])
        self.assertEqual(3,exported["Quantidade solicitada original"])
        self.assertIn("30000001",str(list(book["Histórico"].values)))

import sys
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
import test_purchase_requests as fixtures

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "estoque_app"))
from services import purchase_requests as pr
from purchase_request_export import build_workbook, COLUMNS, XLSX_MIME


def values(sheet):
    return [dict(zip([cell.value for cell in sheet[4]], row))
            for row in sheet.iter_rows(min_row=5, values_only=True)]


class PurchaseRequestExportTests(unittest.TestCase):
    setUpBase = fixtures.PurchaseRequestTests.setUpBase
    setUp = fixtures.PurchaseRequestTests.setUp
    tearDown = fixtures.PurchaseRequestTests.tearDown
    request = fixtures.PurchaseRequestTests.request
    order = fixtures.PurchaseRequestTests.order
    work_order = fixtures.PurchaseRequestTests.work_order

    def export(self, filters=None):
        return load_workbook(build_workbook(pr.export_data(self.db, filters or {}, self.operator),
                                           filters or {}, "OPERADOR"))

    def test_export_includes_all_pages_and_retains_one_row_per_request(self):
        first = self.work_order(3185)
        second = self.work_order(3186)
        original = self.request(work_order_ids=[first, second], sector="PRODUÇÃO")
        base = pr.get(self.db, original["id"])
        copies = [dict(base, id=str(uuid4()), idempotency_key=str(uuid4()), status="CANCELADA")
                  for _ in range(101)]
        self.db.execute(pr.requests.insert(), copies)
        self.db.commit()
        self.assertEqual(100, len(pr.listing(self.db, {})["items"]))
        exported = values(self.export()["Solicitações"])
        self.assertEqual(102, len(exported))
        self.assertEqual(102, len({row["ID da solicitação"] for row in exported}))
        linked = next(row for row in exported if row["ID da solicitação"] == original["id"])
        self.assertIn("TA003185", linked["Referência / O.S. / setor"])
        self.assertIn("TA003186", linked["Referência / O.S. / setor"])

    def test_export_filters_origin_status_dates_and_linked_work_order(self):
        self.request(sector="GERAL")
        order = self.work_order(3191)
        row = self.request(origin="PCP", user=self.pcp, sector="PRODUÇÃO", work_order_ids=[order])
        filters = {"q": "3191", "origin": "PCP", "status": "SOLICITADA",
                   "from": "2026-10-15", "to": "2026-10-15", "page": "999"}
        book = self.export(filters)
        exported = values(book["Solicitações"])
        self.assertEqual([row["id"]], [r["ID da solicitação"] for r in exported])
        history = values(book["Histórico"])
        self.assertTrue(history)
        self.assertEqual({row["id"]}, {r["ID da solicitação"] for r in history})
        self.assertIn("Origem: PCP", book["Solicitações"]["A2"].value)

    def test_empty_export_has_headers_and_valid_sheets(self):
        book = self.export({"q": "inexistente"})
        self.assertEqual(["Solicitações", "Histórico"], book.sheetnames)
        self.assertEqual([], values(book["Solicitações"]))
        self.assertEqual(len(COLUMNS), book["Solicitações"].max_column)
        self.assertEqual("D5", book["Solicitações"].freeze_panes)
        self.assertEqual("A4:AF4", book["Solicitações"].auto_filter.ref)

    def test_quantities_and_dates_are_typed_and_time_is_brasilia(self):
        row = self.request(quantity="2.375")
        stamp = datetime(2026, 10, 8, 2, 30, tzinfo=timezone.utc)
        self.db.execute(pr.requests.update().where(pr.requests.c.id == row["id"])
                        .values(created_at=stamp, updated_at=stamp))
        self.db.commit()
        book = self.export()
        exported = values(book["Solicitações"])[0]
        self.assertEqual(2.375, exported["Quantidade"])
        self.assertEqual(datetime(2026, 10, 7, 23, 30), exported["Solicitada em (Brasília)"])
        self.assertEqual(datetime(2026, 10, 15), exported["Necessidade"])
        self.assertEqual("s", book["Solicitações"]["D5"].data_type)

    def test_user_text_cannot_be_executed_as_formula(self):
        self.request(notes='=HYPERLINK("https://example.invalid","teste")')
        book = self.export()
        cell = book["Solicitações"]["Y5"]
        self.assertEqual("s", cell.data_type)
        self.assertTrue(cell.value.startswith("="))
        self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))

    def test_history_keeps_before_after_actor_and_reason(self):
        row = self.request()
        pr.transition(self.db, row["id"], {"action": "EDITAR", "version": row["version"],
            "sku_codigo": "MAT-001", "quantity": "3", "needed_at": "2026-10-20",
            "sector": "GERAL", "work_order_ids": [], "reason": "Corrigir necessidade"}, self.operator)
        self.db.commit()
        history = values(self.export()["Histórico"])
        changed = next(item for item in history if item["Ação"] == "EDITAR" and item["Campo"] == "quantity")
        self.assertEqual("OPERADOR", changed["Usuário"])
        self.assertEqual("Corrigir necessidade", changed["Motivo"])
        self.assertEqual("2.000", changed["Antes"])
        self.assertEqual("3", changed["Depois"])

    def test_long_audit_values_are_not_silently_truncated(self):
        row = self.request()
        data = pr.export_data(self.db, {})
        text = "á" * 70000
        data["events"] = [{"id": "evento", "request_id": row["id"], "action": "EDIÇÃO",
            "before_data": {}, "after_data": {"text": text}, "reason": "", "actor": "PCP"}]
        book = load_workbook(build_workbook(data))
        exported = values(book["Histórico"])
        self.assertEqual(text, "".join(item["Depois"] for item in exported))
        self.assertEqual([1, 2, 3], [item["Parte"] for item in exported])

    def test_historical_reference_review_and_normalization_are_exported(self):
        row = self.request()
        pr.reference_backfill.create(self.engine, checkfirst=True)
        self.db.execute(pr.requests.update().where(pr.requests.c.id == row["id"])
                        .values(reference="OS 9999"))
        self.db.execute(pr.reference_backfill.insert().values(
            request_id=row["id"], original_reference="OS 9999", result="SEM_MATCH",
            linked_work_orders=0, unresolved_tokens="{9999}"))
        self.db.commit()
        book = self.export()
        exported = values(book["Solicitações"])[0]
        self.assertEqual("SEM_MATCH", exported["Revisão do vínculo histórico"])
        self.assertEqual("9999", exported["Referências históricas sem vínculo"])
        self.assertTrue(any(item["Ação"] == "NORMALIZACAO_REFERENCIAS"
                            for item in values(book["Histórico"])))

    def test_closed_request_displays_completed_without_changing_audit(self):
        row = self.request()
        self.order([row["id"]])
        self.db.commit()
        self.assertEqual("CONCLUÍDO", values(self.export()["Solicitações"])[0]["Classificação"])


class PurchaseRequestExportRouteTests(unittest.TestCase):
    setUpBase = fixtures.PurchaseRequestRouteTests.setUpBase
    setUp = fixtures.PurchaseRequestRouteTests.setUp
    tearDown = fixtures.PurchaseRequestRouteTests.tearDown
    request = fixtures.PurchaseRequestRouteTests.request

    def test_local_and_internal_exports_return_excel_attachment(self):
        self.request()
        for prefix in ("/api/erp/", "/api/erp/internal/"):
            with self.subTest(prefix=prefix):
                response = self.client.get(prefix + "purchase-requests/export.xlsx?q=MAT")
                self.assertEqual(200, response.status_code)
                self.assertEqual(XLSX_MIME, response.mimetype)
                self.assertIn("attachment", response.headers["Content-Disposition"])
                self.assertEqual("no-store", response.headers["Cache-Control"])
                self.assertEqual(1, len(values(load_workbook(BytesIO(response.data))["Solicitações"])))
                response.close()

    def test_inactive_user_cannot_export(self):
        self.actor.active = False
        response = self.client.get("/api/erp/purchase-requests/export.xlsx")
        self.assertEqual(403, response.status_code)
        self.assertEqual("application/json", response.mimetype)

    def test_invalid_date_does_not_download_a_broken_workbook(self):
        response = self.client.get("/api/erp/purchase-requests/export.xlsx?from=ontem")
        self.assertEqual(400, response.status_code)
        self.assertIn("Filtro de data", response.json["error"])


if __name__ == "__main__":
    unittest.main()

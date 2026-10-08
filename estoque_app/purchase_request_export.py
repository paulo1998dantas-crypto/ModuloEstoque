"""Native ERP Excel export. Reads canonical workflow data; no stock movements."""
import json
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BRASILIA = ZoneInfo("America/Sao_Paulo")
STATUS = {"SOLICITADA": "Solicitada", "EM_COMPRAS": "Em compras",
          "CONCLUIDA": "Concluído", "CANCELADA": "Cancelada"}
# Each source request remains one row, even when it has several linked O.S.
COLUMNS = [
    ("Solicitação", "number", 20), ("Origem", "origin_label", 18),
    ("Classificação", "type_label", 20), ("SKU", "sku_codigo", 16),
    ("Descrição", "descricao", 65), ("Unidade", "unidade", 12),
    ("Quantidade", "quantity", 16), ("Necessidade", "needed_at", 16),
    ("Status", "status_label", 18), ("Setor", "sector", 20),
    ("Referência / O.S. / setor", "reference_display", 60),
    ("Referência original", "reference", 40),
    ("Solicitante", "requested_by", 22), ("Solicitada em (Brasília)", "created_at", 25),
    ("Comprador", "buyer", 22), ("Atualizada em (Brasília)", "updated_at", 25),
    ("Pedido existente / criado", "order_number", 24), ("Fornecedor", "supplier", 45),
    ("Status do pedido", "order_status", 25),
    ("Saldo pendente no pedido", "anticipation_pending_quantity", 25),
    ("Previsão do pedido", "anticipation_delivery_date", 22),
    ("Nova data negociada", "anticipation_confirmed_delivery_date", 23),
    ("Conclusão compras (Brasília)", "completed_at", 27),
    ("Concluída por", "completed_by", 22), ("Observações", "notes", 65),
    ("Revisão do vínculo histórico", "reference_review_result", 28),
    ("Referências históricas sem vínculo", "reference_review_tokens", 38),
    ("ID da solicitação", "id", 40), ("IDs das O.S. vinculadas", "work_order_ids", 50),
    ("ID do pedido criado / alocado", "purchase_order_id", 40),
    ("ID do pedido de antecipação", "anticipation_order_id", 40),
    ("Versão", "version", 12),
]
DATE_FIELDS = {"needed_at", "anticipation_delivery_date", "anticipation_confirmed_delivery_date"}
TIME_FIELDS = {"created_at", "updated_at", "completed_at"}
NUMBER_FIELDS = {"quantity", "anticipation_pending_quantity", "version"}


def local_time(value):
    if not value:
        return None
    parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(BRASILIA).replace(tzinfo=None)


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)

def _history_text(field, value):
    if value and field in TIME_FIELDS:
        return local_time(value).strftime("%d/%m/%Y %H:%M:%S")
    if value and field in DATE_FIELDS:
        return date.fromisoformat(str(value)).strftime("%d/%m/%Y")
    return _text(value)


def _write(cell, value, number_format=None):
    cell.value = value
    # User-entered text is always text, never an Excel formula or external link.
    if isinstance(value, str):
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if number_format:
        cell.number_format = number_format


def _sheet(wb, name, title, headers, widths, metadata):
    sheet = wb.create_sheet(name)
    _write(sheet.cell(1, 1), title)
    sheet.cell(1, 1).font = Font(size=16, bold=True, color="163D62")
    _write(sheet.cell(2, 1), metadata)
    sheet.merge_cells("A1:F1")
    sheet.merge_cells("A2:L2")
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 34
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(4, column)
        _write(cell, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163D62")
        sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1]
    sheet.row_dimensions[4].height = 34
    sheet.freeze_panes = "D5"
    sheet.sheet_view.showGridLines = False
    return sheet


def build_workbook(data, filters=None, exported_by="", generated_at=None):
    """Return an in-memory .xlsx with complete rows and their audit events."""
    generated_at = generated_at or datetime.now(timezone.utc)
    stamp = local_time(generated_at).strftime("%d/%m/%Y %H:%M:%S")
    filters = filters or {}
    filter_names = {"q": "Busca", "status": "Status", "origin": "Origem", "from": "Necessidade de", "to": "Até"}
    selected = "; ".join(f"{label}: {filters[key]}" for key, label in filter_names.items() if filters.get(key)) or "Todos"
    metadata = f"Gerado em Brasília: {stamp}. Usuário: {exported_by}. Filtros: {selected}. Solicitações: {len(data['items'])}."
    wb = Workbook()
    wb.remove(wb.active)
    sheet = _sheet(wb, "Solicitações", "Workflow de solicitações de compra",
                   [col[0] for col in COLUMNS], [col[2] for col in COLUMNS], metadata)
    for row_index, raw in enumerate(data["items"], 5):
        row = dict(raw)
        row.update(number="SOL-" + row["id"][:8].upper(),
                   origin_label="Almoxarifado" if row["origin"] == "ESTOQUE" else "PCP",
                   type_label="CONCLUÍDO" if row["status"] == "CONCLUIDA" else
                       "ANTECIPAÇÃO" if row["request_type"] == "ANTECIPACAO" else "NOVA COMPRA",
                   status_label=STATUS.get(row["status"], row["status"]),
                   order_number=row.get("numero_oc") or row.get("anticipation_numero_oc"),
                   supplier=row.get("fornecedor_nome") or row.get("anticipation_fornecedor_nome"),
                   order_status=row.get("purchase_status") or row.get("anticipation_purchase_status"))
        for column, (_, field, _) in enumerate(COLUMNS, 1):
            value, fmt = row.get(field), None
            if field in DATE_FIELDS:
                value = date.fromisoformat(str(value)) if value else None
                fmt = "dd/mm/yyyy"
            elif field in TIME_FIELDS:
                value, fmt = local_time(value), "dd/mm/yyyy hh:mm:ss"
            elif field in NUMBER_FIELDS:
                value = Decimal(str(value)) if value is not None else None
                fmt = "0" if field == "version" else "#,##0.###"
            elif field in {"work_order_ids", "reference_review_tokens"}:
                value = " / ".join(map(str, value or []))
            else:
                value = _text(value)
            _write(sheet.cell(row_index, column), value, fmt)
        sheet.row_dimensions[row_index].height = 48
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(COLUMNS))}{max(4, sheet.max_row)}"

    history = _sheet(wb, "Histórico", "Rastreabilidade das solicitações exportadas",
                     ["Solicitação", "SKU", "Ação", "Usuário", "Data / hora (Brasília)",
                      "Campo", "Antes", "Depois", "Motivo", "ID do evento", "ID da solicitação", "Parte"],
                     [20, 16, 32, 24, 25, 28, 65, 65, 65, 40, 40, 12], metadata)
    by_id = {row["id"]: row for row in data["items"]}
    row_index = 5
    for event in data.get("events", []):
        request_row = by_id.get(event["request_id"])
        if not request_row:
            continue
        before, after = event.get("before_data") or {}, event.get("after_data") or {}
        fields = [key for key in sorted(set(before) | set(after))
                  if key != "idempotency_key" and before.get(key) != after.get(key)] or [""]
        for field in fields:
            old, new = _history_text(field, before.get(field)), _history_text(field, after.get(field))
            # Excel limits a text cell to 32,767 characters. Split without losing audit data.
            for offset in range(0, max(len(old), len(new), 1), 32000):
                values = ["SOL-" + request_row["id"][:8].upper(), request_row["sku_codigo"],
                          event["action"], event.get("actor", ""), local_time(event.get("created_at")),
                          field, old[offset:offset + 32000], new[offset:offset + 32000],
                          event.get("reason", ""), event["id"], event["request_id"], offset // 32000 + 1]
                for column, value in enumerate(values, 1):
                    _write(history.cell(row_index, column), value, "dd/mm/yyyy hh:mm:ss" if column == 5 else None)
                history.row_dimensions[row_index].height = 48
                row_index += 1
    history.auto_filter.ref = f"A4:L{max(4, history.max_row)}"
    result = BytesIO()
    wb.save(result)
    result.seek(0)
    return result
